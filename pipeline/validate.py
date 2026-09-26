"""Data-quality report: what the pipeline checked, what it fixed, and what it could not confirm."""
from __future__ import annotations

import datetime as dt
import math
import re
from collections import Counter

import openpyxl
from rapidfuzz import fuzz

from common import INTERIM, PROCESSED, SAMPLE_XLSX, read_json, write_json


def hav_mi(a, b):
    la1, lo1, la2, lo2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 2 * 3958.8 * math.asin(math.sqrt(h))


def audit_sample(projects, pairs):
    wb = openpyxl.load_workbook(SAMPLE_XLSX, data_only=True)
    rows = list(wb["projects"].iter_rows(values_only=True))
    hdr, prj = rows[0], [dict(zip(rows[0], r)) for r in rows[1:] if r[0]]
    ov_rows = list(wb["overlaps"].iter_rows(values_only=True))
    ovs = [dict(zip(ov_rows[0], r)) for r in ov_rows[1:] if r[0]]
    issues = []
    # 1) distances reproduce?
    byid = {p["project_id"]: p for p in prj}
    repro = []
    for o in ovs:
        a, b = byid[o["project_id_a"]], byid[o["project_id_b"]]
        mi = hav_mi((a["lat_center"], a["lon_center"]), (b["lat_center"], b["lon_center"]))
        repro.append({"overlap_id": o["overlap_id"], "claimed_mi": o["distance_mi"], "recomputed_mi": round(mi, 2),
                      "match": abs(mi - o["distance_mi"]) < 0.01})
    # 2) same substation, different coordinates
    coords = {}
    for p in prj:
        for side in ("a", "b"):
            n, lat, lon = p[f"name_{side}"], p[f"lat_{side}"], p[f"lon_{side}"]
            if n and lat is not None:
                key = re.sub(r"\b(sub|substation)\b", "", str(n).lower()).strip()
                coords.setdefault(key, set()).add((round(lat, 6), round(lon, 6), p["project_id"]))
    for k, v in coords.items():
        pts = {(a, b) for a, b, _ in v}
        if len(pts) > 1:
            (p1, p2) = list(pts)[:2]
            issues.append({"severity": "warn", "check": "inconsistent coordinates",
                           "detail": f"'{k.title()}' has {len(pts)} different coordinates ({hav_mi(p1, p2):.2f} mi apart) across {sorted(x[2] for x in v)}"})
    # 3) unlocated endpoints
    for p in prj:
        for side in ("a", "b"):
            if p[f"name_{side}"] and p[f"lat_{side}"] is None:
                issues.append({"severity": "warn", "check": "unlocated endpoint",
                               "detail": f"{p['project_id']}: '{p[f'name_{side}']}' has no coordinates (center falls back to the other endpoint)"})
    # 4) date formats
    for p in prj:
        v = p["in_service_date"]
        if isinstance(v, (int, float)):
            issues.append({"severity": "info", "check": "date stored as number",
                           "detail": f"{p['project_id']}: in_service_date is the raw Excel serial {int(v)} "
                                     f"(= {(dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()})"})
    # 5) overlap slots
    over = [p["project_id"] for p in prj if (p.get("overlap_count") or 0) > 3]
    issues.append({"severity": "info", "check": "schema limit",
                   "detail": "projects sheet stores at most 3 overlaps per project (overlap_1..3); "
                             f"our full run finds projects with up to {max(Counter([x['a'] for x in pairs if x['tier']]).values(), default=0)} nearby Georgia projects - exported as a long table instead"})
    # 6) do we find their pairs?
    def match(name, utility):
        cands = [p for p in projects if (p["utility"] == "DESC") == (utility.startswith("Dominion"))]
        best = max(cands, key=lambda p: fuzz.token_set_ratio(name.lower(), p["name"].lower()))
        return best["id"]
    found = []
    pair_set = {(x["a"], x["b"]): x for x in pairs}
    for o in ovs:
        a, b = match(o["project_name_a"], o["utility_a"]), match(o["project_name_b"], o["utility_b"])
        x = pair_set.get((a, b))
        found.append({"overlap_id": o["overlap_id"], "ours": bool(x and x["tier"]), "our_tier": x and x["tier"],
                      "our_mi": x and x["mi"], "their_center_mi": o["distance_mi"], "pair_id": x and x["id"],
                      "their_gap_days": o.get("time_gap (day)")})
    return {"reproduced": repro, "issues": issues, "found": found, "n_projects": len(prj), "n_overlaps": len(ovs)}


def main():
    projects = read_json(PROCESSED / "projects.json")
    pairs = read_json(PROCESSED / "pairs.json")
    churn = read_json(INTERIM / "gpc_churn.json")
    desc = [p for p in projects if p["utility"] == "DESC"]
    ga = [p for p in projects if p["utility"] != "DESC"]
    source_issues = []
    for p in desc:
        for n in p.get("notes", []):
            source_issues.append({"project": p["id"], "name": p["name"], "detail": n, "page": p["source_page"]})
        if p.get("cost_total") and p["cost_total"] < 2_000_000:
            source_issues.append({"project": p["id"], "name": p["name"], "page": p["source_page"],
                                  "detail": f"total ${p['cost_total']:,} is below the list's own '$2M and above' threshold"})
        if p.get("in_service") and p["in_service"] > "2028-12-31":
            source_issues.append({"project": p["id"], "name": p["name"], "page": p["source_page"],
                                  "detail": f"in-service {p['in_service']} falls outside the list's 2024-2028 window"})
        if "transformer" in (p.get("need") or "").lower() and "transformer" not in (p["description"] + p["name"]).lower() \
                and "autobank" not in p["name"].lower():
            source_issues.append({"project": p["id"], "name": p["name"], "page": p["source_page"],
                                  "detail": "need statement mentions transformers but the scope does not (likely copy-paste in the filing)"})
    dup = Counter(p["name"] for p in desc)
    for n, c in dup.items():
        if c > 1:
            source_issues.append({"project": "-", "name": n, "page": None,
                                  "detail": f"{c} different projects share this exact name; keyed by Project ID instead"})
    changes = Counter()
    for p in ga:
        note = (p.get("change_note") or "").lower()
        changes["new" if "new" in note else "delayed" if "delay" in note else "advanced" if "advanc" in note
                else "no change" if "no change" in note else "other"] += 1
    gpc_new_miles = sum(p["miles"] or 0 for p in ga if p["work_type"] == "new_line")
    report = {
        "stages": [
            {"stage": "Extract DESC PDF", "in": "44 pages", "out": f"{len(desc)} projects", "ok": len(desc) == 44},
            {"stage": "Extract GA ITS plan", "in": "Table 2 + 208 detail pages",
             "out": f"{len(ga)} projects ({churn['table2_count']} Table 2 rows joined)", "ok": len(ga) == 208},
            {"stage": "Locate endpoints", "in": f"{len(projects)} projects",
             "out": f"{sum(p['located'] for p in projects)} placed on map", "ok": True},
            {"stage": "Compare pairs", "in": f"{sum(p['located'] for p in desc)} x {sum(p['located'] for p in ga)}",
             "out": f"{sum(1 for x in pairs if x['tier'])} within 40 km", "ok": True},
        ],
        "confidence": {u: dict(Counter(p["confidence"] for p in projects if (p["utility"] == "DESC") == (u == "DESC")))
                       for u in ("DESC", "GA")},
        "source_issues": source_issues,
        "plan_changes": dict(changes),
        "cancelled": churn["cancelled"], "completed": churn["completed"],
        "table1_check": {"filed_new_row_miles_10yr": 1142, "extracted_new_line_miles": round(gpc_new_miles, 1),
                         "note": "Table 1 totals new-ROW line miles for the whole GA ITS; extracted miles come only from detail pages that state a length"},
        "sample_audit": audit_sample(projects, pairs),
        "unlocated": [{"id": p["id"], "name": p["name"], "endpoints": [e["name"] for e in p["endpoints"]]}
                      for p in projects if not p["located"]],
    }
    write_json(PROCESSED / "quality.json", report)
    sa = report["sample_audit"]
    print("sample reproduced:", sum(r["match"] for r in sa["reproduced"]), "/", len(sa["reproduced"]))
    print("sample pairs we also find:", sum(f["ours"] for f in sa["found"]), "/", len(sa["found"]), sa["found"])
    print("sample issues:", len(sa["issues"]))
    for i in sa["issues"]:
        print("  -", i["check"], ":", i["detail"])
    print("source issues:", len(source_issues))
    for s in source_issues:
        print("  -", s["name"][:50], ":", s["detail"])
    print("plan changes:", report["plan_changes"], "| table1:", report["table1_check"])


if __name__ == "__main__":
    main()
